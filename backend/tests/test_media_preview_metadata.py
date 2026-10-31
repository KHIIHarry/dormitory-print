"""Production media-preview contract: source bytes, normalized PDF, range and auth.

These tests use the real PostgreSQL fixtures, the real upload route, the real
asset worker, and a real LibreOffice conversion for DOCX. They exist because the
mobile UI showed "resource unavailable" for an original DOCX preview and
"0.0 MB" for a small-but-valid DOCX: the frontend must consume the server's
normalized PDF URL and byte count, not infer either from the local File object.

No production behaviour is changed here; this locks the already-correct
backend contract so the frontend guidance in
``docs/qa/2026-10-31-media-preview-handoff.md`` has executable evidence.
"""
from __future__ import annotations

import io
import secrets
from urllib.parse import unquote

import pytest

pymupdf = pytest.importorskip('pymupdf')

from test_media import guest, make_pdf, process_asset, seed_config, upload

APP_MODULES = ['media', 'growth', 'mail', 'auth']


def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading('DormPrint 原文件预览验收', 0)
    document.add_paragraph('第一页：移动端原始文档预览。')
    document.add_page_break()
    document.add_paragraph('第二页：服务端归一化 PDF。')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _headers(client):
    return {
        'Origin': 'http://testserver',
        'Idempotency-Key': secrets.token_urlsafe(24),
        'X-CSRF-Token': unquote(client.cookies.get('csrfToken', '')),
    }


def _json_get(client, path):
    response = client.get(path, headers=_headers(client))
    assert response.status_code == 200, response.text
    return response.json()['data']


def test_docx_upload_reports_real_bytes_and_exposes_normalized_pdf(
    client, client_factory, db
):
    seed_config(db)
    guest(client)
    payload = _docx_bytes()
    assert len(payload) > 0

    created = upload(
        client,
        payload,
        'preview-contract.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    assert created.status_code == 200, created.text
    result = created.json()['data']
    asset_id = result['assetId']

    # The upload result is bytes, never MB and never a rounded display value.
    assert result['size'] == len(payload)
    assert result['size'] > 0
    assert result['url'] == f'/api/v1/assets/{asset_id}/content'
    # A DOCX cannot have a synchronous preview; the worker produces it.
    assert result['status'] == 'QUEUED'
    assert result['normalizedUrl'] is None

    queued = _json_get(client, f'/api/v1/assets/{asset_id}')
    assert queued['size'] == len(payload)
    assert queued['kind'] == 'DOCX'
    assert queued['normalizedUrl'] is None

    process_asset(db, asset_id)
    detail = _json_get(client, f'/api/v1/assets/{asset_id}')

    assert detail['status'] == 'READY'
    assert detail['size'] == len(payload)
    assert detail['pages'] == 2
    assert detail['normalizedAssetId']
    assert detail['normalizedUrl'] == (
        f"/api/v1/assets/{detail['normalizedAssetId']}/content"
    )
    assert detail['previewUrl'] == detail['normalizedUrl']
    assert detail['mime'].startswith(
        'application/vnd.openxmlformats-officedocument.wordprocessingml'
    )

    # Draft recovery and the original-preview panel need the same byte count.
    listed = _json_get(client, '/api/v1/assets?purpose=PRINT_SOURCE')
    row = next(item for item in listed['items'] if item['assetId'] == asset_id)
    assert row['size'] == len(payload)

    raw = client.get(detail['url'])
    assert raw.status_code == 200
    assert raw.content == payload
    assert raw.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.wordprocessingml'
    )
    assert raw.headers['content-disposition'] == 'inline'
    assert raw.headers['x-content-type-options'] == 'nosniff'
    assert raw.headers['accept-ranges'] == 'bytes'

    preview = client.get(detail['normalizedUrl'])
    assert preview.status_code == 200
    assert preview.headers['content-type'] == 'application/pdf'
    assert preview.headers['content-disposition'] == 'inline'
    assert preview.headers['accept-ranges'] == 'bytes'
    assert preview.headers['x-content-type-options'] == 'nosniff'
    assert preview.content.startswith(b'%PDF-')

    normalized = pymupdf.open(stream=preview.content, filetype='pdf')
    try:
        assert normalized.page_count == detail['pages'] == 2
    finally:
        normalized.close()

    # The normalized preview is still an authorized, private asset.
    other = guest(client_factory())
    assert other.get(detail['normalizedUrl']).status_code == 404
    assert other.get(detail['url']).status_code == 404


def test_pdf_upload_size_and_normalized_preview_are_byte_exact(
    client, client_factory, db
):
    seed_config(db)
    guest(client)
    payload = make_pdf(2)

    created = upload(client, payload, 'preview-contract.pdf')
    assert created.status_code == 200, created.text
    result = created.json()['data']
    asset_id = result['assetId']
    assert result['size'] == len(payload)

    process_asset(db, asset_id)
    detail = _json_get(client, f'/api/v1/assets/{asset_id}')
    assert detail['status'] == 'READY'
    assert detail['size'] == len(payload)
    assert detail['pages'] == 2
    assert detail['normalizedUrl'] == (
        f"/api/v1/assets/{detail['normalizedAssetId']}/content"
    )

    raw = client.get(detail['url'])
    assert raw.status_code == 200
    assert raw.content == payload
    assert raw.headers['content-type'].startswith('application/pdf')

    preview = client.get(detail['normalizedUrl'])
    assert preview.status_code == 200
    assert preview.headers['content-type'] == 'application/pdf'
    assert preview.content.startswith(b'%PDF-')

    # Mobile Safari and pdf.js both use range requests; keep them byte-exact.
    ranged = client.get(
        detail['normalizedUrl'], headers={'Range': 'bytes=0-9'}
    )
    assert ranged.status_code == 206
    assert ranged.content == preview.content[:10]
    assert ranged.headers['content-range'] == (
        f'bytes 0-9/{len(preview.content)}'
    )

    other = guest(client_factory())
    assert other.get(detail['normalizedUrl']).status_code == 404
