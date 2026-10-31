"""Real document conversion, pagination, photo geometry and publishing tests.

These tests build real PDF/DOCX/PPTX/XLSX/JPEG/PNG/WebP/HEIC files and run the
same LibreOffice/PyMuPDF/Pillow pipeline used in production. They never accept
mocks or synthetic page counts as evidence.
"""
from __future__ import annotations

import hashlib
import io
import json
import shutil
import zipfile
from pathlib import Path

import pytest

pymupdf = pytest.importorskip('pymupdf')


def _make_pdf(path: Path, pages):
    """pages: list of (width_pt, height_pt, rotation_degrees)."""
    doc = pymupdf.open()
    for width, height, rotation in pages:
        page = doc.new_page(width=width, height=height)
        page.draw_rect(
            pymupdf.Rect(0, 0, width, height),
            color=(0.9, 0.1, 0.1),
            fill=(0.9, 0.1, 0.1),
            width=1,
        )
        page.insert_text((24, 40), f'page {page.number + 1}', fontsize=14)
        if rotation:
            page.set_rotation(rotation)
    doc.save(path)
    return path


def _render_bbox(pdf_path: Path, page_index: int = 0, dpi: int = 72):
    """Bounding box of non-white pixels in a rendered page, in points."""
    doc = pymupdf.open(pdf_path)
    page = doc[page_index]
    pix = page.get_pixmap(dpi=dpi)
    scale = 72.0 / dpi
    xs, ys = [], []
    samples = pix.samples
    stride = pix.stride
    channels = pix.n
    for y in range(pix.height):
        row = y * stride
        for x in range(pix.width):
            offset = row + x * channels
            if min(samples[offset:offset + 3]) < 245:
                xs.append(x)
                ys.append(y)
    if not xs:
        return None
    return (
        min(xs) * scale,
        min(ys) * scale,
        (max(xs) + 1) * scale,
        (max(ys) + 1) * scale,
    )


def test_non_a4_pages_fit_inside_a4_without_cropping(tmp_path):
    from app.pdf_engine import A4_HEIGHT_PT, A4_WIDTH_PT, normalize_pdf

    source = _make_pdf(
        tmp_path / 'mixed.pdf',
        [(595.276, 841.89, 0), (841.89, 595.276, 0), (1190.55, 1683.78, 0)],
    )
    out = tmp_path / 'normalized.pdf'
    result = normalize_pdf(source, out, max_pages=20)

    assert result.pages == 3
    doc = pymupdf.open(out)
    assert doc.page_count == 3
    for page in doc:
        assert page.rect.width == pytest.approx(A4_WIDTH_PT, abs=0.5)
        assert page.rect.height == pytest.approx(A4_HEIGHT_PT, abs=0.5)
    # The landscape A4 page is letterboxed, not cropped: the full-page red
    # rectangle keeps its true 841.89/595.276 aspect ratio inside A4.
    landscape = result.page_geometry[1]
    assert landscape['adapted'] is True
    assert landscape['scale'] == pytest.approx(A4_WIDTH_PT / 841.89, rel=1e-3)
    bbox = _render_bbox(out, 1, dpi=72)
    assert bbox is not None
    width, height = bbox[2] - bbox[0], bbox[3] - bbox[1]
    assert width / height == pytest.approx(841.89 / 595.276, rel=0.01)
    assert width <= A4_WIDTH_PT + 1
    # Content must not touch/cross the A4 edge (letterboxing leaves margins).
    assert bbox[0] >= -0.5 and bbox[1] >= -0.5
    # One rendered pixel may fall on the page boundary at 72dpi.
    assert bbox[2] <= A4_WIDTH_PT + 1.5
    assert bbox[3] <= A4_HEIGHT_PT + 1.5
    # The A3 page is scaled down, never cut: width fills A4.
    a3 = result.page_geometry[2]
    assert a3['scale'] == pytest.approx(0.5, abs=0.02)
    assert _render_bbox(out, 2, dpi=72)[2] <= A4_WIDTH_PT + 2


def test_duplex_blank_backs_and_page_sheet_map_for_seven_pages_three_copies(tmp_path):
    from app.pdf_engine import build_document_artifact, normalize_pdf

    source = _make_pdf(tmp_path / 'seven.pdf', [(595.276, 841.89, 0)] * 7)
    normalized = tmp_path / 'seven-a4.pdf'
    assert normalize_pdf(source, normalized, max_pages=50).pages == 7

    art = build_document_artifact(
        normalized,
        tmp_path / 'artifact.pdf',
        segment_id='seg-1',
        pages=[1, 2, 3, 4, 5, 6, 7],
        sides='DUPLEX',
        copies=3,
        color_mode='COLOR',
        flip='LONG_EDGE',
    )
    assert (art.pages, art.sheets, art.blank_backs) == (24, 12, 3)
    doc = pymupdf.open(tmp_path / 'artifact.pdf')
    assert doc.page_count == 24
    # Copy 1 = pages 1..8 with page 8 the system blank back of sheet 4.
    assert len(art.page_map) == 24
    first = art.page_map[:8]
    assert [entry['sourcePageIndex'] for entry in first[:7]] == [1, 2, 3, 4, 5, 6, 7]
    assert first[7]['sourcePageIndex'] is None and first[7]['blank'] is True
    assert first[7]['side'] == 'BACK' and first[7]['sheetIndex'] == 4
    assert first[0]['side'] == 'FRONT' and first[1]['side'] == 'BACK'
    assert first[1]['sheetIndex'] == first[0]['sheetIndex'] == 1
    assert [entry['copyIndex'] for entry in art.page_map] == [1] * 8 + [2] * 8 + [3] * 8
    assert [entry['sheetOrdinal'] for entry in art.page_map[:8]] == [1, 1, 2, 2, 3, 3, 4, 4]
    assert first[0]['sheetId'] == 'c1-s1'
    assert art.page_map[-1]['sheetId'] == 'c3-s4'


def test_single_sided_and_segment_order(tmp_path):
    from app.pdf_engine import build_document_artifact, normalize_pdf

    source = _make_pdf(tmp_path / 'order.pdf', [(595.276, 841.89, 0)] * 7)
    normalized = tmp_path / 'order-a4.pdf'
    normalize_pdf(source, normalized, max_pages=50)
    art = build_document_artifact(
        normalized,
        tmp_path / 'ordered.pdf',
        segment_id='seg-2',
        pages=[7, 2, 5],
        sides='SINGLE',
        copies=2,
        color_mode='BW',
        flip='NONE',
    )
    assert (art.pages, art.sheets, art.blank_backs) == (6, 6, 0)
    assert [entry['sourcePageIndex'] for entry in art.page_map] == [7, 2, 5, 7, 2, 5]
    assert all(entry['side'] == 'FRONT' for entry in art.page_map)
    with pytest.raises(Exception):
        build_document_artifact(
            normalized, tmp_path / 'bad.pdf', segment_id='x', pages=[1, 1],
            sides='SINGLE', copies=1, color_mode='BW', flip='NONE',
        )
    with pytest.raises(Exception):
        build_document_artifact(
            normalized, tmp_path / 'bad2.pdf', segment_id='x', pages=[1, 8],
            sides='SINGLE', copies=1, color_mode='BW', flip='NONE',
        )


def test_document_artifact_insets_content_by_configured_a4_margins(tmp_path):
    from app.pdf_engine import (
        A4_HEIGHT_PT,
        A4_WIDTH_PT,
        EngineError,
        build_document_artifact,
    )

    mm = 72.0 / 25.4
    source = _make_pdf(
        tmp_path / 'edge.pdf', [(A4_WIDTH_PT, A4_HEIGHT_PT, 0)])
    full = build_document_artifact(
        source,
        tmp_path / 'full-bleed.pdf',
        segment_id='seg-full',
        pages=[1],
        sides='SINGLE',
        copies=1,
        color_mode='COLOR',
        flip='NONE',
    )
    assert full.pages == 1
    full_bbox = _render_bbox(tmp_path / 'full-bleed.pdf')
    assert full_bbox is not None
    assert full_bbox[0] <= 1 and full_bbox[1] <= 1

    margins = {'top': 10, 'right': 10, 'bottom': 10, 'left': 10}
    inset = build_document_artifact(
        source,
        tmp_path / 'inset.pdf',
        segment_id='seg-inset',
        pages=[1],
        sides='SINGLE',
        copies=1,
        color_mode='COLOR',
        flip='NONE',
        margins_mm=margins,
    )
    assert inset.pages == 1
    doc = pymupdf.open(tmp_path / 'inset.pdf')
    assert doc[0].rect.width == pytest.approx(A4_WIDTH_PT, abs=0.5)
    assert doc[0].rect.height == pytest.approx(A4_HEIGHT_PT, abs=0.5)
    bbox = _render_bbox(tmp_path / 'inset.pdf')
    assert bbox is not None
    assert bbox[0] >= 10 * mm - 1
    assert bbox[1] >= 10 * mm - 1
    assert bbox[2] <= (210 - 10) * mm + 1
    assert bbox[3] <= (297 - 10) * mm + 1
    # Equal margins change the aspect ratio, so the page is width-limited and
    # its full printable width is used.
    assert bbox[2] - bbox[0] == pytest.approx(190 * mm, abs=1.5)

    with pytest.raises(EngineError) as exc:
        build_document_artifact(
            source,
            tmp_path / 'invalid-margins.pdf',
            segment_id='seg-bad',
            pages=[1],
            sides='SINGLE',
            copies=1,
            color_mode='COLOR',
            flip='NONE',
            margins_mm={'top': 10, 'right': 10, 'bottom': 10, 'left': 200},
        )
    assert exc.value.code == 'PAPER_GEOMETRY_INVALID'


def test_document_artifact_bakes_annotations_when_insetting(tmp_path):
    from app.pdf_engine import A4_HEIGHT_PT, A4_WIDTH_PT, build_document_artifact

    source = tmp_path / 'annotated.pdf'
    doc = pymupdf.open()
    page = doc.new_page(width=A4_WIDTH_PT, height=A4_HEIGHT_PT)
    page.insert_text((50, 50), 'body')
    page.add_freetext_annot(pymupdf.Rect(100, 100, 300, 150), 'STAMP')
    doc.save(source)
    doc.close()

    build_document_artifact(
        source,
        tmp_path / 'annotated-inset.pdf',
        segment_id='seg-annot',
        pages=[1],
        sides='SINGLE',
        copies=1,
        color_mode='COLOR',
        flip='NONE',
        margins_mm={'top': 10, 'right': 10, 'bottom': 10, 'left': 10},
    )

    out = pymupdf.open(tmp_path / 'annotated-inset.pdf')
    assert 'STAMP' in out[0].get_text()
    assert list(out[0].annots()) == []
    out.close()


def test_page_allocation_requires_every_page_printed_or_excluded():
    from app.pdf_engine import EngineError, validate_page_allocation

    validate_page_allocation(5, [[1, 2], [3, 4]], [5])
    for ranges, excluded in (
        ([[1, 2], [2, 3]], [4, 5]),
        ([[1, 2], [3]], [5]),
        ([[1, 2], [3, 4]], [4, 5]),
    ):
        with pytest.raises(EngineError) as exc:
            validate_page_allocation(5, ranges, excluded)
        assert exc.value.code == 'PAGE_ALLOCATION_INVALID'
    for ranges, excluded in (([[1, 2], [3, 4]], [0, 5]), ([[1, 2], [3, 4]], [6])):
        with pytest.raises(EngineError) as exc:
            validate_page_allocation(5, ranges, excluded)
        assert exc.value.code == 'PAGE_RANGE_INVALID'


def test_real_docx_pptx_xlsx_and_legacy_xls_conversion(tmp_path):
    from app.pdf_engine import normalize_pdf, resolve_soffice, source_to_normalized_pdf

    if not resolve_soffice():
        pytest.skip('LibreOffice is not installed in this environment')

    docx = tmp_path / 'real.docx'
    import docx as docx_module

    document = docx_module.Document()
    document.add_heading('DormPrint 转换验收', 0)
    document.add_paragraph('第一页真实段落。')
    document.add_page_break()
    document.add_paragraph('第二页真实段落。')
    document.save(docx)

    pptx = tmp_path / 'real.pptx'
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for title in ('第一张幻灯片', '第二张幻灯片'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text_frame.text = '真实转换内容'
    presentation.save(pptx)

    xlsx = tmp_path / 'real.xlsx'
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '打印区域'
    for row in range(1, 21):
        for column in range(1, 6):
            sheet.cell(row=row, column=column, value=f'{column}-{row}')
    sheet.print_area = 'A1:E10'
    workbook.save(xlsx)

    for source, expected in ((docx, 2), (pptx, 2), (xlsx, 1)):
        pdf, meta = source_to_normalized_pdf(
            source,
            source.suffix.lstrip('.').upper(),
            tmp_path / (source.stem + '-work'),
            max_pages=30,
            excel_options=None,
        )
        probe = pymupdf.open(pdf)
        assert probe.page_count == expected, (source.name, probe.page_count)
        assert meta['converter'] in ('libreoffice', 'native-pdf')
        assert meta['normalizedPages'] == expected

    # A real legacy .xls produced by LibreOffice must also round-trip.
    xls = tmp_path / 'real.xls'
    import subprocess

    subprocess.run(
        [
            resolve_soffice(), '--headless', '--norestore',
            f'-env:UserInstallation=file://{tmp_path / "xls-profile"}',
            '--convert-to', 'xls', '--outdir', str(tmp_path), str(xlsx),
        ],
        check=True, capture_output=True, timeout=120,
    )
    if xls.exists():
        pdf, meta = source_to_normalized_pdf(
            xls,
            'XLS',
            tmp_path / 'xls-work',
            max_pages=30,
            excel_options={'printArea': 'A1:E10', 'printTitlesRows': '1:1'},
        )
        assert pymupdf.open(pdf).page_count == 1
        assert meta['excel']['printArea'] == 'A1:E10'


def test_excel_options_are_strict_and_safe():
    from app.pdf_engine import EngineError, validate_excel_options

    options = validate_excel_options(
        {
            'sheetName': '打印区域',
            'printArea': 'A1:F40',
            'printTitlesRows': '1:2',
            'printTitlesColumns': 'A:B',
            'orientation': 'LANDSCAPE',
            'paperSize': 'A4',
            'fitToWidth': 1,
            'fitToHeight': 0,
            'scalePercent': 90,
            'gridlines': True,
            'marginsMm': {'top': 10, 'right': 8, 'bottom': 10, 'left': 8},
        }
    )
    assert options.fit_to_width == 1 and options.orientation == 'LANDSCAPE'

    bad = [
        {'printArea': 'A0:F40'},
        {'printArea': 'A1:F40:G50'},
        {'printArea': 'XFE1:XFE2'},
        {'printArea': 'A1:ZZ99999'},
        {'unknownKey': 1},
        {'orientation': 'SIDEWAYS'},
        {'paperSize': 'A3'},
        {'fitToWidth': 999},
        {'scalePercent': 5},
        {'marginsMm': {'left': 100, 'right': 100}},
        {'printTitlesRows': '0:3'},
        {'printArea': 'A1:E200000'},
    ]
    for payload in bad:
        with pytest.raises(EngineError) as exc:
            validate_excel_options(payload)
        assert exc.value.code == 'EXCEL_OPTIONS_INVALID'


def test_ooxml_macros_external_links_and_mismatched_extension_are_rejected(tmp_path):
    import docx as docx_module
    from app.pdf_engine import EngineError, inspect_upload

    document = docx_module.Document()
    document.add_paragraph('safe body')
    source = tmp_path / 'safe.docx'
    document.save(source)

    # A docx renamed to xlsx must be rejected from real zip contents, not suffix.
    renamed = tmp_path / 'renamed.xlsx'
    shutil.copy(source, renamed)
    with pytest.raises(EngineError) as exc:
        inspect_upload(renamed, renamed.name, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert exc.value.code in ('MIME_MISMATCH', 'INVALID_OFFICE')

    macro = tmp_path / 'macro.docx'
    with zipfile.ZipFile(source) as origin, zipfile.ZipFile(macro, 'w') as target:
        for info in origin.infolist():
            target.writestr(info, origin.read(info.filename))
        target.writestr('word/vbaProject.bin', b'malicious')
    with pytest.raises(EngineError) as exc:
        inspect_upload(macro, macro.name, '')
    assert exc.value.code == 'OFFICE_MACROS_FORBIDDEN'

    linked = tmp_path / 'linked.docx'
    with zipfile.ZipFile(source) as origin, zipfile.ZipFile(linked, 'w') as target:
        for info in origin.infolist():
            if info.filename == 'word/_rels/document.xml.rels':
                continue
            target.writestr(info, origin.read(info.filename))
        target.writestr(
            'word/_rels/document.xml.rels',
            b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="file:///etc/passwd" TargetMode="External"/></Relationships>',
        )
    with pytest.raises(EngineError) as exc:
        inspect_upload(linked, linked.name, '')
    assert exc.value.code == 'OFFICE_EXTERNAL_CONTENT_FORBIDDEN'


def test_pdf_signature_mismatch_and_encrypted_pdf_are_rejected(tmp_path):
    from app.pdf_engine import EngineError, inspect_upload, validate_pdf

    png = tmp_path / 'image.png'
    from PIL import Image

    Image.new('RGB', (32, 32), 'white').save(png)
    with pytest.raises(EngineError) as exc:
        inspect_upload(png, 'fake.pdf', 'application/pdf')
    assert exc.value.code in ('MIME_MISMATCH', 'INVALID_FILE')

    pdf = _make_pdf(tmp_path / 'ok.pdf', [(595.276, 841.89, 0)])
    assert inspect_upload(pdf, 'ok.pdf', 'application/pdf').kind == 'PDF'

    encrypted = tmp_path / 'locked.pdf'
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((20, 20), 'secret')
    doc.save(
        encrypted,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw='owner',
        user_pw='user',
    )
    with pytest.raises(EngineError) as exc:
        validate_pdf(encrypted, max_pages=10)
    assert exc.value.code == 'ENCRYPTED_PDF'


def test_heic_real_decode_exif_and_contain_millimetre_geometry(tmp_path):
    from PIL import Image
    from pillow_heif import register_heif_opener

    from app.pdf_engine import PaperSpec, render_photo_pdf

    register_heif_opener()
    heic = tmp_path / 'photo.heic'
    Image.new('RGB', (400, 300), (220, 30, 30)).save(heic, format='HEIF')

    out = tmp_path / 'photo.pdf'
    render = render_photo_pdf(
        heic,
        out,
        copies=2,
        orientation='LANDSCAPE',
        fit='CONTAIN',
        rotation_degrees=0,
        crop=None,
        paper=PaperSpec(127.0, 89.0, {'top': 3.0, 'right': 3.0, 'bottom': 3.0, 'left': 3.0}),
        dpi=150,
    )
    doc = pymupdf.open(out)
    assert doc.page_count == 2
    assert doc[0].rect.width == pytest.approx(127 / 25.4 * 72, abs=0.3)
    assert doc[0].rect.height == pytest.approx(89 / 25.4 * 72, abs=0.3)
    assert render.print_width_mm == pytest.approx(121.0)
    assert render.print_height_mm == pytest.approx(83.0)
    assert render.source_width == 400 and render.source_height == 300
    assert render.copies == 2
    bbox = _render_bbox(out, 0, dpi=150)
    # CONTAIN: 4:3 image in 121x83 print area -> height 83mm, width 110.67mm.
    left_mm, top_mm = bbox[0] / 72 * 25.4, bbox[1] / 72 * 25.4
    right_mm, bottom_mm = bbox[2] / 72 * 25.4, bbox[3] / 72 * 25.4
    assert top_mm == pytest.approx(3.0, abs=0.8)
    assert bottom_mm == pytest.approx(86.0, abs=0.8)
    assert right_mm - left_mm == pytest.approx(83.0 * 4 / 3, abs=0.8)
    assert left_mm > 3.0 + 2 and right_mm < 124.0 - 2

    # EXIF orientation 6 is a real 90° clockwise correction: 40x20 -> 20x40.
    jpeg = tmp_path / 'exif.jpg'
    image = Image.new('RGB', (40, 20), (20, 20, 200))
    exif = image.getexif()
    exif[274] = 6
    image.save(jpeg, exif=exif)
    portrait = render_photo_pdf(
        jpeg,
        tmp_path / 'exif.pdf',
        copies=1,
        orientation='PORTRAIT',
        fit='CONTAIN',
        rotation_degrees=0,
        crop=None,
        paper=PaperSpec(89.0, 127.0, {'top': 3.0, 'right': 3.0, 'bottom': 3.0, 'left': 3.0}),
        dpi=150,
    )
    assert (portrait.source_width, portrait.source_height) == (20, 40)
    assert portrait.exif_orientation == 6

    # Explicit user rotation is applied clockwise after EXIF correction, and
    # crop/paper geometry is computed on the rotated result.
    rotated = render_photo_pdf(
        jpeg,
        tmp_path / 'rotated.pdf',
        copies=1,
        orientation='PORTRAIT',
        fit='CONTAIN',
        rotation_degrees=90,
        crop=None,
        paper=PaperSpec(89.0, 127.0, {'top': 3.0, 'right': 3.0, 'bottom': 3.0, 'left': 3.0}),
        dpi=150,
    )
    # EXIF 6 -> 20x40, then +90° -> 40x20.
    assert (rotated.source_width, rotated.source_height) == (40, 20)
    rotated_bbox = _render_bbox(tmp_path / 'rotated.pdf', 0, dpi=150)
    rotated_width = rotated_bbox[2] - rotated_bbox[0]
    rotated_height = rotated_bbox[3] - rotated_bbox[1]
    # JPEG edge ringing slightly widens the rendered detected box.
    assert rotated_width / rotated_height == pytest.approx(2.0, rel=0.1)


def test_cover_crop_fills_print_area_and_rejects_wrong_aspect(tmp_path):
    from PIL import Image

    from app.pdf_engine import EngineError, PaperSpec, photo_crop, render_photo_pdf

    image_path = tmp_path / 'cover.png'
    Image.new('RGB', (400, 300), (30, 180, 30)).save(image_path)
    paper = PaperSpec(127.0, 89.0, {'top': 3.0, 'right': 3.0, 'bottom': 3.0, 'left': 3.0})
    crop = photo_crop(400, 300, 121.0, 83.0, 1.2, 50, 50)
    render = render_photo_pdf(
        image_path,
        tmp_path / 'cover.pdf',
        copies=1,
        orientation='LANDSCAPE',
        fit='COVER',
        rotation_degrees=0,
        crop=crop,
        paper=paper,
        dpi=150,
    )
    bbox = _render_bbox(tmp_path / 'cover.pdf', 0, dpi=150)
    left_mm, top_mm = bbox[0] / 72 * 25.4, bbox[1] / 72 * 25.4
    right_mm, bottom_mm = bbox[2] / 72 * 25.4, bbox[3] / 72 * 25.4
    assert left_mm == pytest.approx(3.0, abs=0.8)
    assert top_mm == pytest.approx(3.0, abs=0.8)
    assert right_mm == pytest.approx(124.0, abs=0.8)
    assert bottom_mm == pytest.approx(86.0, abs=0.8)
    assert render.crop['width'] * 400 / (render.crop['height'] * 300) == pytest.approx(
        121.0 / 83.0, rel=0.015
    )

    wrong = dict(crop)
    wrong['width'] = wrong['width'] * 0.5
    with pytest.raises(EngineError) as exc:
        render_photo_pdf(
            image_path,
            tmp_path / 'wrong.pdf',
            copies=1,
            orientation='LANDSCAPE',
            fit='COVER',
            rotation_degrees=0,
            crop=wrong,
            paper=paper,
        )
    assert exc.value.code == 'PHOTO_GEOMETRY_INVALID'
    # CONTAIN never accepts a crop: it must not silently crop the user's photo.
    with pytest.raises(EngineError):
        render_photo_pdf(
            image_path,
            tmp_path / 'contain-crop.pdf',
            copies=1,
            orientation='LANDSCAPE',
            fit='CONTAIN',
            rotation_degrees=0,
            crop=crop,
            paper=paper,
        )


def test_bw_grayscale_keeps_text_vector_and_fallback_warns(tmp_path):
    from app.pdf_engine import convert_pdf_to_grayscale

    base = _make_pdf(tmp_path / 'color-base.pdf', [(595.276, 841.89, 0)])
    source = tmp_path / 'color.pdf'
    doc = pymupdf.open(base)
    page = doc[0]
    page.draw_rect(pymupdf.Rect(50, 100, 200, 200), color=(0, 0.5, 0.9), fill=(0, 0.5, 0.9))
    doc.save(source)
    doc.close()

    gray = tmp_path / 'gray-gs.pdf'
    result = convert_pdf_to_grayscale(source, gray, gs_path=None)
    assert result.method == 'ghostscript'
    assert result.rasterized is False
    assert result.warning is None
    gray_doc = pymupdf.open(gray)
    text = gray_doc[0].get_text()
    assert 'page 1' in text
    pix = gray_doc[0].get_pixmap(dpi=72)
    for x, y in ((60, 120), (150, 150), (300, 400)):
        pixel = pix.pixel(x, y)
        assert pixel[0] == pixel[1] == pixel[2]

    fallback = tmp_path / 'gray-raster.pdf'
    import app.pdf_engine as engine

    original_resolver = engine.resolve_ghostscript
    engine.resolve_ghostscript = lambda configured=None: None
    try:
        result = convert_pdf_to_grayscale(source, fallback, dpi=150)
    finally:
        engine.resolve_ghostscript = original_resolver
    assert result.method == 'pymupdf-raster'
    assert result.rasterized is True
    assert result.warning and result.warning['code'] == 'BW_GRAYSCALE_RASTERIZED'
    raster_doc = pymupdf.open(fallback)
    assert raster_doc[0].get_text().strip() == ''
    pix = raster_doc[0].get_pixmap(dpi=72)
    assert pix.pixel(60, 120)[0] == pix.pixel(60, 120)[1]


def test_atomic_publish_verifies_hash_and_preserves_existing_file(tmp_path):
    from app.pdf_engine import EngineError, publish_file

    payload = b'%PDF-1.7 real artifact bytes'
    digest = hashlib.sha256(payload).hexdigest()
    temp = tmp_path / 'temp.part'
    final = tmp_path / 'published.pdf'
    temp.write_bytes(payload)
    assert publish_file(temp, final, digest) == digest
    assert final.read_bytes() == payload

    existing = final.read_bytes()
    duplicate = tmp_path / 'duplicate.part'
    duplicate.write_bytes(payload)
    publish_file(duplicate, final, digest)
    assert final.read_bytes() == existing

    corrupt = tmp_path / 'corrupt.part'
    corrupt.write_bytes(b'not the expected bytes')
    with pytest.raises(EngineError) as exc:
        publish_file(corrupt, tmp_path / 'new.pdf', digest)
    assert exc.value.code == 'ARTIFACT_VERIFY_FAILED'
    assert not (tmp_path / 'new.pdf').exists()
    assert final.read_bytes() == existing


def test_sheet_reprint_preserves_front_back_and_blank_back(tmp_path):
    from app.pdf_engine import (
        build_document_artifact,
        build_sheet_reprint_pdf,
        normalize_pdf,
    )

    source = _make_pdf(tmp_path / 'reprint-src.pdf', [(595.276, 841.89, 0)] * 7)
    normalized = tmp_path / 'reprint-a4.pdf'
    normalize_pdf(source, normalized, max_pages=20)
    artifact = build_document_artifact(
        normalized,
        tmp_path / 'reprint-artifact.pdf',
        segment_id='seg-1',
        pages=[1, 2, 3, 4, 5, 6, 7],
        sides='DUPLEX',
        copies=2,
        color_mode='COLOR',
        flip='LONG_EDGE',
    )
    reprint = build_sheet_reprint_pdf(
        tmp_path / 'reprint-artifact.pdf',
        artifact.page_map,
        ['c2-s1', 'c2-s4'],
        tmp_path / 'reprint.pdf',
    )
    # A duplex physical sheet always reprints both sides: c2-s1 is pages 1/2,
    # c2-s4 is source page 7 plus its genuine blank back.
    assert reprint.pages == 4 and reprint.sheets == 2
    assert list(reprint.selected_sheets) == ['c2-s1', 'c2-s4']
    doc = pymupdf.open(tmp_path / 'reprint.pdf')
    assert doc.page_count == 4
    assert 'page 1' in doc[0].get_text()
    assert 'page 2' in doc[1].get_text()
    assert 'page 7' in doc[2].get_text()
    assert doc[3].get_text().strip() == ''


def test_legacy_doc_ppt_and_webp_are_really_supported(tmp_path):
    import docx as docx_module
    from pptx import Presentation

    from app.pdf_engine import (
        PaperSpec,
        inspect_upload,
        render_photo_pdf,
        resolve_soffice,
        source_to_normalized_pdf,
    )

    if not resolve_soffice():
        pytest.skip('LibreOffice is not installed in this environment')
    soffice = resolve_soffice()

    document = docx_module.Document()
    document.add_paragraph('旧版 DOC 真实转换')
    document.add_page_break()
    document.add_paragraph('第二页')
    docx_path = tmp_path / 'legacy.docx'
    document.save(docx_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = '旧版 PPT 真实转换'
    pptx_path = tmp_path / 'legacy.pptx'
    presentation.save(pptx_path)

    import subprocess

    subprocess.run(
        [soffice, '--headless', '--norestore',
         f'-env:UserInstallation=file://{tmp_path / "legacy-doc-profile"}',
         '--convert-to', 'doc', '--outdir', str(tmp_path), str(docx_path)],
        check=True, capture_output=True, timeout=120,
    )
    subprocess.run(
        [soffice, '--headless', '--norestore',
         f'-env:UserInstallation=file://{tmp_path / "legacy-ppt-profile"}',
         '--convert-to', 'ppt', '--outdir', str(tmp_path), str(pptx_path)],
        check=True, capture_output=True, timeout=120,
    )
    doc_path, ppt_path = tmp_path / 'legacy.doc', tmp_path / 'legacy.ppt'
    assert doc_path.exists() and ppt_path.exists()
    assert inspect_upload(doc_path, 'legacy.doc', 'application/msword').kind == 'DOC'
    assert inspect_upload(ppt_path, 'legacy.ppt', 'application/vnd.ms-powerpoint').kind == 'PPT'

    # Production stores blobs under extensionless content-addressed paths.
    extensionless = tmp_path / ('d' * 40)
    shutil.copy(doc_path, extensionless)
    pdf, meta = source_to_normalized_pdf(
        extensionless, 'DOC', tmp_path / 'doc-work', max_pages=20, excel_options=None
    )
    assert meta['converter'] == 'libreoffice'
    assert pymupdf.open(pdf).page_count == 2
    pdf, meta = source_to_normalized_pdf(
        ppt_path, 'PPT', tmp_path / 'ppt-work', max_pages=20, excel_options=None
    )
    assert pymupdf.open(pdf).page_count == 1

    # WebP is a first-class source image, decoded by Pillow and rendered into
    # the calibrated photo paper geometry.
    from PIL import Image

    webp = tmp_path / 'photo.webp'
    Image.new('RGB', (320, 240), (10, 150, 60)).save(webp, format='WEBP')
    info = inspect_upload(webp, 'photo.webp', 'image/webp')
    assert info.kind == 'WEBP'
    render = render_photo_pdf(
        webp,
        tmp_path / 'webp.pdf',
        copies=1,
        orientation='LANDSCAPE',
        fit='CONTAIN',
        rotation_degrees=0,
        crop=None,
        paper=PaperSpec(127.0, 89.0, {'top': 3.0, 'right': 3.0, 'bottom': 3.0, 'left': 3.0}),
        dpi=150,
    )
    assert render.pages == 1
    assert pymupdf.open(tmp_path / 'webp.pdf').page_count == 1


def test_excel_pagination_options_change_real_page_count(tmp_path):
    import openpyxl

    from app.pdf_engine import source_to_normalized_pdf

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '分页'
    for row in range(1, 121):
        for column in range(1, 6):
            sheet.cell(row=row, column=column, value=f'{column}-{row}')
    source = tmp_path / 'pagination.xlsx'
    workbook.save(source)

    multi, meta = source_to_normalized_pdf(
        source, 'XLSX', tmp_path / 'multi-work', max_pages=100,
        excel_options={
            'sheetName': '分页', 'printArea': 'A1:E120',
            'orientation': 'LANDSCAPE', 'fitToWidth': 1, 'fitToHeight': 0,
            'scalePercent': 100,
        },
    )
    multi_pages = pymupdf.open(multi).page_count
    assert multi_pages > 1
    assert meta['excel']['fitToHeight'] == 0

    single, meta = source_to_normalized_pdf(
        source, 'XLSX', tmp_path / 'single-work', max_pages=100,
        excel_options={
            'sheetName': '分页', 'printArea': 'A1:E120',
            'orientation': 'LANDSCAPE', 'fitToWidth': 1, 'fitToHeight': 1,
        },
    )
    single_pages = pymupdf.open(single).page_count
    assert single_pages == 1
    assert single_pages < multi_pages


def test_office_subprocess_env_excludes_secrets_and_database(tmp_path, monkeypatch):
    import docx as docx_module

    from app import pdf_engine

    document = docx_module.Document()
    document.add_paragraph('conversion environment isolation')
    source = tmp_path / 'env.docx'
    document.save(source)
    monkeypatch.setenv('DORMPRINT_DATABASE_URL', 'postgresql://secret-db/dormprint')
    monkeypatch.setenv('DORMPRINT_SECRET_KEY', 's' * 48)
    captured = {}
    original = pdf_engine._run

    def capture(command, *, timeout, env, cwd):
        captured['command'] = list(command)
        captured['env'] = dict(env)
        return original(command, timeout=timeout, env=env, cwd=cwd)

    monkeypatch.setattr(pdf_engine, '_run', capture)
    output = pdf_engine.office_to_pdf(source, tmp_path / 'env-work')
    assert output.exists()
    env = captured['env']
    assert not any(key.upper().startswith('DORMPRINT') for key in env)
    flattened = ' '.join(str(value) for value in env.values())
    assert 'secret-db' not in flattened
    assert 's' * 48 not in flattened
    assert captured['command'][0].endswith('soffice')


def test_version_probe_subprocess_env_excludes_secrets(monkeypatch):
    import subprocess

    from app import pdf_engine

    captured = []

    def fake_run(argv, **kwargs):
        captured.append(kwargs.get('env'))
        return subprocess.CompletedProcess(argv, 0, stdout=b'probe-version\n', stderr=b'')

    monkeypatch.setenv('DORMPRINT_DATABASE_URL', 'postgresql://secret-db/dormprint')
    monkeypatch.setenv('DORMPRINT_SECRET_KEY', 'k' * 48)
    monkeypatch.setattr(pdf_engine.subprocess, 'run', fake_run)
    pdf_engine._soffice_version_cached.cache_clear()
    pdf_engine._gs_version_cached.cache_clear()
    try:
        assert pdf_engine._soffice_version_cached('/bin/echo') == 'probe-version'
        assert pdf_engine._gs_version_cached('/bin/echo') == 'probe-version'
    finally:
        pdf_engine._soffice_version_cached.cache_clear()
        pdf_engine._gs_version_cached.cache_clear()
    assert captured
    for env in captured:
        assert env is not None
        assert not any(key.upper().startswith('DORMPRINT') for key in env)
        assert 'secret-db' not in ' '.join(str(value) for value in env.values())
